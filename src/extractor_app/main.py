import argparse
import json
import os
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
import pdfplumber
import fitz
from PIL import Image
from transformers import pipeline
import camelot
import easyocr
import docx

SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff"}
SUPPORTED_DOC_EXTENSIONS = {".pdf", ".docx"}


@dataclass
class ChartResult:
    image_id: str
    caption: Optional[str]
    chart_data: Optional[List[Dict[str, Any]]]
    raw_chart_text: Optional[str]
    source_path: str


@dataclass
class TableResult:
    source: str
    page_start: int
    page_end: int
    columns: List[str]
    rows: List[Dict[str, Any]]


@dataclass
class TextResult:
    source: str
    text: str


@dataclass
class ImageResult:
    image_id: str
    source: str
    caption: Optional[str]
    path: str


class ExtractorPipeline:
    def __init__(self) -> None:
        self.chart_pipeline = pipeline(
            "image-to-text",
            model="google/deplot",
            max_new_tokens=512,
        )
        self.caption_pipeline = pipeline(
            "image-to-text",
            model="Salesforce/blip-image-captioning-base",
            max_new_tokens=64,
        )
        self.ocr_reader = easyocr.Reader(["en"], gpu=False)

    def extract_chart(self, image: Image.Image, image_id: str, source_path: str) -> ChartResult:
        chart_output = self.chart_pipeline(image)[0]["generated_text"]
        chart_data = parse_deplot_output(chart_output)
        caption = self.caption_pipeline(image)[0]["generated_text"]
        return ChartResult(
            image_id=image_id,
            caption=caption,
            chart_data=chart_data,
            raw_chart_text=chart_output,
            source_path=source_path,
        )

    def extract_caption(self, image: Image.Image) -> Optional[str]:
        return self.caption_pipeline(image)[0]["generated_text"]

    def extract_text(self, image: Image.Image, source: str) -> TextResult:
        ocr_results = self.ocr_reader.readtext(image, detail=0)
        return TextResult(source=source, text="\n".join(ocr_results))


def parse_deplot_output(raw_text: str) -> Optional[List[Dict[str, Any]]]:
    if not raw_text:
        return None
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    if not lines:
        return None
    delimiter = "|" if "|" in lines[0] else "\t"
    rows = [line.strip("|") for line in lines]
    table = [row.split(delimiter) for row in rows]
    headers = [header.strip() for header in table[0]]
    data_rows = []
    for row in table[1:]:
        normalized = [cell.strip() for cell in row]
        row_dict = {headers[idx]: normalized[idx] if idx < len(normalized) else "" for idx in range(len(headers))}
        data_rows.append(row_dict)
    return data_rows


def extract_tables_from_pdf(pdf_path: Path) -> List[TableResult]:
    tables: List[TableResult] = []
    camelot_tables = camelot.read_pdf(str(pdf_path), pages="all", flavor="stream")
    for table in camelot_tables:
        df = table.df
        headers = df.iloc[0].tolist()
        rows = []
        for _, row in df.iloc[1:].iterrows():
            rows.append({headers[idx]: row.iloc[idx] for idx in range(len(headers))})
        tables.append(
            TableResult(
                source=str(pdf_path),
                page_start=table.page,
                page_end=table.page,
                columns=headers,
                rows=rows,
            )
        )
    return merge_tables(tables)


def merge_tables(tables: List[TableResult]) -> List[TableResult]:
    if not tables:
        return tables
    merged: List[TableResult] = [tables[0]]
    for table in tables[1:]:
        last = merged[-1]
        if table.columns == last.columns and table.page_start == last.page_end + 1:
            last.rows.extend(table.rows)
            last.page_end = table.page_end
        else:
            merged.append(table)
    return merged


def extract_images_from_pdf(pdf_path: Path, output_dir: Path) -> List[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    image_paths: List[Path] = []
    with fitz.open(pdf_path) as doc:
        for page_index in range(len(doc)):
            page = doc.load_page(page_index)
            for image_index, image in enumerate(page.get_images(full=True)):
                xref = image[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                ext = base_image["ext"]
                image_path = output_dir / f"page_{page_index + 1}_img_{image_index + 1}.{ext}"
                image_path.write_bytes(image_bytes)
                image_paths.append(image_path)
    return image_paths


def extract_from_docx(docx_path: Path) -> Dict[str, Any]:
    doc = docx.Document(docx_path)
    tables: List[TableResult] = []
    for table in doc.tables:
        rows = []
        headers = [cell.text for cell in table.rows[0].cells]
        for row in table.rows[1:]:
            rows.append({headers[idx]: cell.text for idx, cell in enumerate(row.cells)})
        tables.append(
            TableResult(
                source=str(docx_path),
                page_start=1,
                page_end=1,
                columns=headers,
                rows=rows,
            )
        )
    texts = [para.text for para in doc.paragraphs if para.text.strip()]
    return {
        "tables": tables,
        "texts": [TextResult(source=str(docx_path), text="\n".join(texts))] if texts else [],
        "images": [],
    }


def process_file(input_path: Path, output_dir: Path) -> Dict[str, Any]:
    pipeline = ExtractorPipeline()
    charts: List[ChartResult] = []
    tables: List[TableResult] = []
    texts: List[TextResult] = []
    images: List[ImageResult] = []

    if input_path.suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS:
        image = Image.open(input_path)
        chart = pipeline.extract_chart(image, image_id=input_path.stem, source_path=str(input_path))
        charts.append(chart)
        texts.append(pipeline.extract_text(image, source=str(input_path)))
    elif input_path.suffix.lower() == ".pdf":
        tables.extend(extract_tables_from_pdf(input_path))
        image_paths = extract_images_from_pdf(input_path, output_dir / "images")
        for image_path in image_paths:
            image = Image.open(image_path)
            chart = pipeline.extract_chart(image, image_id=image_path.stem, source_path=str(input_path))
            charts.append(chart)
            caption = chart.caption
            images.append(
                ImageResult(
                    image_id=image_path.stem,
                    source=str(input_path),
                    caption=caption,
                    path=str(image_path),
                )
            )
            texts.append(pipeline.extract_text(image, source=str(image_path)))
        with pdfplumber.open(input_path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    texts.append(TextResult(source=f"{input_path}#page-{page_index}", text=page_text))
    elif input_path.suffix.lower() == ".docx":
        docx_data = extract_from_docx(input_path)
        tables.extend(docx_data["tables"])
        texts.extend(docx_data["texts"])
    else:
        raise ValueError(f"Unsupported file type: {input_path.suffix}")

    return {
        "input": str(input_path),
        "charts": charts,
        "tables": tables,
        "texts": texts,
        "images": images,
    }


def write_outputs(result: Dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "extraction.json"
    with json_path.open("w", encoding="utf-8") as handle:
        json.dump(
            {
                "input": result["input"],
                "charts": [asdict(chart) for chart in result["charts"]],
                "tables": [asdict(table) for table in result["tables"]],
                "texts": [asdict(text) for text in result["texts"]],
                "images": [asdict(image) for image in result["images"]],
            },
            handle,
            indent=2,
            ensure_ascii=False,
        )

    with pd.ExcelWriter(output_dir / "extraction.xlsx") as writer:
        if result["charts"]:
            chart_rows = []
            for chart in result["charts"]:
                if chart.chart_data:
                    for row in chart.chart_data:
                        chart_rows.append({"image_id": chart.image_id, **row})
                else:
                    chart_rows.append({"image_id": chart.image_id, "raw": chart.raw_chart_text or ""})
            pd.DataFrame(chart_rows).to_excel(writer, sheet_name="charts", index=False)
        if result["tables"]:
            table_rows = []
            for table in result["tables"]:
                for row in table.rows:
                    table_rows.append({"source": table.source, **row})
            pd.DataFrame(table_rows).to_excel(writer, sheet_name="tables", index=False)
        if result["texts"]:
            pd.DataFrame([asdict(text) for text in result["texts"]]).to_excel(
                writer, sheet_name="texts", index=False
            )
        if result["images"]:
            pd.DataFrame([asdict(image) for image in result["images"]]).to_excel(
                writer, sheet_name="images", index=False
            )


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract charts, tables, and text into JSON/XLSX.")
    parser.add_argument("input", help="Path to PDF, image, or DOCX")
    parser.add_argument("--output-dir", default="output", help="Directory to store results")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    output_dir = Path(args.output_dir)
    result = process_file(input_path, output_dir)
    write_outputs(result, output_dir)


if __name__ == "__main__":
    main()
